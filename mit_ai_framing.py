"""
MIT AI News: how language shapes the image of AI
Header (title + summary) vs Body  ->  Positive / Neutral-Cautious / Risk-oriented
Outputs: scored data, category summary, distinctive vocabulary, contrast word clouds, themes.
Only needs pandas, numpy, scikit-learn, matplotlib, Pillow, openpyxl.
"""
import re, math, random, sys
from collections import Counter
import numpy as np, pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from PIL import Image, ImageDraw, ImageFont
from sklearn.feature_extraction.text import TfidfVectorizer, ENGLISH_STOP_WORDS
from sklearn.decomposition import NMF

IN  = sys.argv[1] if len(sys.argv) > 1 else "MIT_AI_ARTICLES.csv"
OUT = sys.argv[2] if len(sys.argv) > 2 else "."
random.seed(1); np.random.seed(1)

# ------------------------------------------------------------------ lexicons
# A trailing * means "any word starting with this stem".
POSITIVE = """improv* accura* efficien* faster fast speed* enabl* breakthrough* novel innovat* advanc*
powerful success* outperform* state-of-the-art better robust effective* effectiv* promis* boost* enhanc*
accelerat* scalab* discover* capabilit* achiev* reliab* transform* revolution* unlock* empower* streamlin*
cutting-edge opportunit* benefit* excit* impressive best strong* superior optimiz* optimis* generaliz*
versatile flexible sophisticated groundbreaking pioneer* milestone""".split()

CAUTIOUS = """could may might possibl* uncertain* limited limitation* challeng* caution* careful* tradeoff*
trade-off* depends unclear remain* although despite whether potential* potentially likely unlikely perhaps
hurdle* difficult* nuanc* ambigu* inconsisten* tentative preliminary early-stage""".split()

RISK = """risk* bias* biased harm* danger* unsafe safety threat* vulnerab* privacy surveillance misinformation
disinformation deepfake* hallucinat* discriminat* unfair* fairness inequalit* inequit* misuse malicious
unintended adversarial attack* breach* manipulat* exploit* fraud* cheat* plagiaris* ethic* governance
regulat* policymak* accountab* transparen* trustworth* liabilit* copyright concern* worr* fear* warn*
jailbreak* toxic* misleading mislead* flaw* fail* failure* error* mistake* undermin* disrupt* displac*
job-loss unemploy* vulnerable erode* abuse* illegal unethical safeguard* oversight""".split()

LEX = {"Positive": POSITIVE, "Cautious": CAUTIOUS, "Risk": RISK}

def compile_lex(words):
    stems  = tuple(w[:-1] for w in words if w.endswith("*"))
    exact  = {w for w in words if not w.endswith("*")}
    return stems, exact

COMPILED = {k: compile_lex(v) for k, v in LEX.items()}

def match(tok, key):
    stems, exact = COMPILED[key]
    return tok in exact or tok.startswith(stems)

# ------------------------------------------------------------------ text prep
TOKEN = re.compile(r"[a-z][a-z\-']*[a-z]|[a-z]")
BOILER = set("""mit said says say like new using use used uses also one two three work works working research
researchers researcher study studies paper team professor university institute department just
according well way ways make makes made many much even may first however while would could
ghassemi autor jacob carlone jameel lincoln wang rus agrawal""".split())
STOP = set(ENGLISH_STOP_WORDS) | BOILER

def tokens(text):
    return TOKEN.findall(str(text).lower().replace("’", "'"))

def content_tokens(text):
    return [t for t in tokens(text) if t not in STOP and len(t) > 2 and not t.endswith("'s")]

# ------------------------------------------------------------------ scoring
def score(text):
    toks = tokens(text)
    n = max(len(toks), 1)
    hits = {k: sum(match(t, k) for t in toks) for k in LEX}
    return n, hits

# Classification rule (fixed before looking at results, then checked for sensitivity):
#   Positive words are ~4x more frequent than risk words in ordinary science news, so raw counts
#   would call almost everything "positive". Each rate is therefore divided by its corpus mean
#   (an index of 1.0 = typical MIT AI article), separately for header and body.
#   net = (pos_idx - risk_idx) / (pos_idx + risk_idx)      in [-1, 1]
#   Positive         : net >=  T      (clearly leans on capability / progress vocabulary)
#   Risk-oriented    : net <= -T      (clearly leans on challenge / bias / safety / governance vocabulary)
#   Neutral/Cautious : otherwise      (balanced, hedged, or no evaluative wording at all)
def classify(pos_idx, risk_idx, T=0.30):
    tot = pos_idx + risk_idx
    if tot == 0:
        return "Neutral/Cautious", 0.0
    net = (pos_idx - risk_idx) / tot
    if net >= T:
        return "Positive", net
    if net <= -T:
        return "Risk-oriented", net
    return "Neutral/Cautious", net

CATS = ["Positive", "Neutral/Cautious", "Risk-oriented"]

df = pd.read_csv(IN)
df = df.drop(columns=[c for c in df.columns if c.startswith("Unnamed")])
df["header"] = df["title"].str.strip() + ". " + df["summary"].str.strip()
df["year"] = df["datetime"].str[:4]

rows = []
for part in ["header", "body"]:
    for i, txt in df[part].items():
        n, h = score(txt)
        rows.append((i, part, n, h["Positive"], h["Cautious"], h["Risk"]))
sc = pd.DataFrame(rows, columns=["idx", "part", "n_tokens", "pos", "caut", "risk"])
for part in ["header", "body"]:
    s = sc[sc.part == part].set_index("idx")
    df[f"{part}_tokens"] = s.n_tokens
    for c in ["pos", "caut", "risk"]:
        df[f"{part}_{c}"] = s[c]
        df[f"{part}_{c}_per1k"] = (s[c] / s.n_tokens * 1000).round(2)
    pos_rate, risk_rate = s.pos / s.n_tokens, s.risk / s.n_tokens
    df[f"{part}_pos_idx"]  = (pos_rate  / pos_rate.mean()).round(3)
    df[f"{part}_risk_idx"] = (risk_rate / risk_rate.mean()).round(3)
    res = [classify(p, r) for p, r in zip(df[f"{part}_pos_idx"], df[f"{part}_risk_idx"])]
    df[f"{part}_label"] = [r[0] for r in res]
    df[f"{part}_net"] = [round(r[1], 3) for r in res]

df["consistent"] = df.header_label == df.body_label

# sensitivity of the body/header label shares to the threshold
sens = []
for T in [0.2, 0.3, 0.4, 0.5]:
    for part in ["header", "body"]:
        s = pd.Series([classify(p, r, T)[0] for p, r in zip(df[f"{part}_pos_idx"], df[f"{part}_risk_idx"])])
        sh = s.value_counts(normalize=True).reindex(CATS).fillna(0).round(3)
        sens.append({"threshold": T, "part": part, **sh.to_dict()})
sens = pd.DataFrame(sens)

# ------------------------------------------------------------------ category summary
summ = []
for part in ["header", "body"]:
    vc = df[f"{part}_label"].value_counts().reindex(CATS).fillna(0).astype(int)
    for c in CATS:
        summ.append({"text": part, "category": c, "n_articles": vc[c], "share": round(vc[c] / len(df), 3)})
summ = pd.DataFrame(summ)

ct = pd.crosstab(df.header_label, df.body_label).reindex(index=CATS, columns=CATS).fillna(0).astype(int)
ct.index.name = "header \\ body"

rates = pd.DataFrame({
    part: {k: df[f"{part}_{k}_per1k"].mean().round(2) for k in ["pos", "caut", "risk"]}
    for part in ["header", "body"]
}).rename(index={"pos": "Positive per 1k words", "caut": "Cautious per 1k words", "risk": "Risk per 1k words"})

# ------------------------------------------------------------------ distinctive vocabulary (weighted log-odds, Monroe et al.)
def distinctive(part, min_count=15, top=40):
    docs = {c: Counter() for c in CATS}
    dfreq = {c: Counter() for c in CATS}          # in how many articles of the category a term appears
    for lab, txt in zip(df[f"{part}_label"], df[part]):
        toks = content_tokens(txt)
        docs[lab].update(toks); dfreq[lab].update(set(toks))
    tot = sum(docs.values(), Counter())
    vocab = [w for w, c in tot.items() if c >= min_count] if part == "body" else [w for w, c in tot.items() if c >= 3]
    a = {w: 0.05 * tot[w] for w in vocab}
    a0 = sum(a.values())
    out = {}
    for c in CATS:
        rest = Counter()
        for o in CATS:
            if o != c:
                rest.update(docs[o])
        n1, n2 = sum(docs[c].values()), sum(rest.values())
        recs = []
        for w in vocab:
            y1, y2 = docs[c][w], rest[w]
            d = (math.log((y1 + a[w]) / (n1 + a0 - y1 - a[w])) -
                 math.log((y2 + a[w]) / (n2 + a0 - y2 - a[w])))
            z = d / math.sqrt(1 / (y1 + a[w]) + 1 / (y2 + a[w]))
            seed = any(match(w, k) for k in LEX)
            if dfreq[c][w] >= (4 if part == "body" else 2):
                recs.append((w, y1, y2, round(z, 2), seed))
        r = pd.DataFrame(recs, columns=["term", "count_in_category", "count_elsewhere", "z_score", "is_lexicon_word"])
        out[c] = r.sort_values("z_score", ascending=False).head(top).reset_index(drop=True)
    return out

dist_body, dist_head = distinctive("body"), distinctive("header")

# ------------------------------------------------------------------ word cloud (self-contained; no wordcloud package)
FONT_LATIN = font_manager.findfont(font_manager.FontProperties(family="DejaVu Sans", weight="bold"))
PALETTES = {
    "Positive":         ["#1b7f4b", "#2a9d62", "#14613a", "#3fae7a"],
    "Neutral/Cautious": ["#b07a00", "#c98f10", "#8a6100", "#d6a436"],
    "Risk-oriented":    ["#b3261e", "#cc3b32", "#8c1d17", "#d9574e"],
}

def word_cloud(weights, palette, W=900, H=600, max_words=55, fmin=16, fmax=95):
    items = sorted(weights.items(), key=lambda kv: -kv[1])[:max_words]
    wmax, wmin = items[0][1], items[-1][1]
    occ = np.zeros((H, W), dtype=bool)
    img = Image.new("RGB", (W, H), "white")
    drw = ImageDraw.Draw(img)
    for rank, (word, w) in enumerate(items):
        t = (w - wmin) / (wmax - wmin + 1e-9)
        size = int(fmin + (fmax - fmin) * t ** 0.8)
        font = ImageFont.truetype(FONT_LATIN, size)
        l, tp, r, b = drw.textbbox((0, 0), word, font=font, stroke_width=3)
        tw, th = r - l, b - tp
        m = Image.new("L", (tw, th), 0)
        ImageDraw.Draw(m).text((-l, -tp), word, font=font, fill=255, stroke_width=3, stroke_fill=255)
        mask = np.array(m) > 0
        placed = False
        ang0 = random.uniform(0, 2 * math.pi)
        for step in range(6000):
            ang = ang0 + step * 0.21
            rad = 2.0 + step * 0.16
            x = int(W / 2 + rad * 1.5 * math.cos(ang) - tw / 2)
            y = int(H / 2 + rad * math.sin(ang) - th / 2)
            if x < 2 or y < 2 or x + tw > W - 2 or y + th > H - 2:
                continue
            if not (occ[y:y + th, x:x + tw] & mask).any():
                occ[y:y + th, x:x + tw] |= mask
                drw.text((x - l, y - tp), word, font=font, fill=palette[rank % len(palette)])
                placed = True
                break
    return img

def cloud_figure(dist, title, fname):
    fig, axes = plt.subplots(1, 3, figsize=(18, 5.6))
    labels = {c: c for c in CATS}
    for ax, c in zip(axes, CATS):
        d = dist[c]
        w = {t: z for t, z in zip(d.term, d.z_score) if z > 0}
        ax.imshow(word_cloud(w, PALETTES[c]))
        ax.set_title(labels[c], fontsize=15, color=PALETTES[c][0], fontweight="bold")
        ax.axis("off")
    fig.suptitle(title, fontsize=16, y=1.02)
    fig.tight_layout()
    fig.savefig(f"{OUT}/{fname}", dpi=150, bbox_inches="tight")
    plt.close(fig)

cloud_figure(dist_body, "Contrast word clouds - Article Body (size = distinctiveness vs. the other two categories)", "fig2_wordclouds_body.png")
cloud_figure(dist_head, "Contrast word clouds - Article Header (title + summary)", "fig3_wordclouds_header.png")

# ------------------------------------------------------------------ themes (NMF on body TF-IDF)
K = 8
vec = TfidfVectorizer(tokenizer=content_tokens, lowercase=False, token_pattern=None, min_df=8, max_df=0.4, ngram_range=(1, 1))
X = vec.fit_transform(df["body"])
nmf = NMF(n_components=K, init="nndsvda", random_state=1, max_iter=600)
Wm = nmf.fit_transform(X)
terms = np.array(vec.get_feature_names_out())
theme_terms = {k: terms[np.argsort(-nmf.components_[k])[:12]].tolist() for k in range(K)}
df["theme_id"] = Wm.argmax(axis=1)

# ------------------------------------------------------------------ theme labels + theme x category
THEME_LABELS = {
    0: "LLMs, language & reasoning",
    1: "Education, students & society",
    2: "Biology, proteins & genomics",
    3: "Materials & chemistry design",
    4: "Robotics & control",
    5: "Vision, image & video generation",
    6: "Health & clinical care",
    7: "Energy, climate & infrastructure",
}
df["theme"] = df.theme_id.map(THEME_LABELS)
theme_rows = []
for k, lab in THEME_LABELS.items():
    sub = df[df.theme_id == k]
    vc = sub.body_label.value_counts(normalize=True).reindex(CATS).fillna(0)
    theme_rows.append({"theme": lab, "n_articles": len(sub), "top_terms": ", ".join(theme_terms[k][:10]),
                       "Positive %": round(100 * vc["Positive"], 1),
                       "Neutral/Cautious %": round(100 * vc["Neutral/Cautious"], 1),
                       "Risk-oriented %": round(100 * vc["Risk-oriented"], 1),
                       "avg risk per 1k words": round(sub.body_risk_per1k.mean(), 2),
                       "avg positive per 1k words": round(sub.body_pos_per1k.mean(), 2)})
themes = pd.DataFrame(theme_rows).sort_values("Risk-oriented %", ascending=False)

# ------------------------------------------------------------------ figures 1 and 4
COL = {"Positive": "#2a9d62", "Neutral/Cautious": "#d6a436", "Risk-oriented": "#cc3b32"}
fig, ax = plt.subplots(1, 2, figsize=(13, 4.6), gridspec_kw={"width_ratios": [1.1, 1]})
share = summ.pivot(index="text", columns="category", values="share").reindex(["header", "body"])[CATS]
left = np.zeros(2)
for c in CATS:
    ax[0].barh(["Header\n(title+summary)", "Body"], share[c].values, left=left, color=COL[c], label=c)
    for i, v in enumerate(share[c].values):
        if v > 0.04:
            ax[0].text(left[i] + v / 2, i, f"{v:.0%}", ha="center", va="center", color="white", fontweight="bold")
    left += share[c].values
ax[0].invert_yaxis(); ax[0].set_xlim(0, 1); ax[0].set_title("Share of 314 articles by framing"); ax[0].legend(loc="lower center", bbox_to_anchor=(0.5, -0.28), ncol=3, frameon=False)
ax[0].xaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1))
x = np.arange(3); wdt = 0.36
vals_h = [rates.loc[r, "header"] for r in rates.index]; vals_b = [rates.loc[r, "body"] for r in rates.index]
ax[1].bar(x - wdt / 2, vals_h, wdt, label="Header", color="#1f3b5c"); ax[1].bar(x + wdt / 2, vals_b, wdt, label="Body", color="#8aa6c6")
ax[1].set_xticks(x); ax[1].set_xticklabels(["Positive", "Cautious", "Risk"]); ax[1].set_ylabel("Lexicon hits per 1,000 words")
ax[1].set_title("Vocabulary density"); ax[1].legend(frameon=False)
for i, (h, b) in enumerate(zip(vals_h, vals_b)):
    ax[1].text(i - wdt / 2, h + 0.4, f"{h:.1f}", ha="center", fontsize=9); ax[1].text(i + wdt / 2, b + 0.4, f"{b:.1f}", ha="center", fontsize=9)
for a_ in ax: a_.spines[["top", "right"]].set_visible(False)
fig.tight_layout(); fig.savefig(f"{OUT}/fig1_sentiment_overview.png", dpi=150, bbox_inches="tight"); plt.close(fig)

fig, ax = plt.subplots(figsize=(10, 5))
t = themes.sort_values("Risk-oriented %")
bottom = np.zeros(len(t))
for c, col in zip(CATS, ["Positive %", "Neutral/Cautious %", "Risk-oriented %"]):
    ax.barh(t.theme + "  (n=" + t.n_articles.astype(str) + ")", t[col], left=bottom, color=COL[c], label=c)
    bottom += t[col].values
ax.set_xlim(0, 100); ax.set_xlabel("% of articles in theme (body framing)")
ax.set_title("Themes x framing: where does risk-oriented language concentrate?")
ax.legend(loc="lower center", bbox_to_anchor=(0.5, -0.25), ncol=3, frameon=False)
ax.spines[["top", "right"]].set_visible(False)
fig.tight_layout(); fig.savefig(f"{OUT}/fig4_themes_by_framing.png", dpi=150, bbox_inches="tight"); plt.close(fig)

# ------------------------------------------------------------------ examples + lexicon sheets
def examples(part, lab, asc, n=8):
    d = df[df[f"{part}_label"] == lab].sort_values(f"{part}_net", ascending=asc).head(n)
    d = d[["title", "publication_date", f"{part}_pos", f"{part}_risk", f"{part}_net", "url"]]
    d.columns = ["title", "publication_date", "pos_hits", "risk_hits", "net_index", "url"]
    return d.assign(text_part=part, category=lab)
ex = pd.concat([examples("body", "Risk-oriented", True), examples("body", "Positive", False),
                examples("header", "Risk-oriented", True), examples("header", "Positive", False)])
lex = pd.DataFrame({"category": [k for k, v in LEX.items() for _ in v], "term (* = stem)": [w for v in LEX.values() for w in v]})

# ------------------------------------------------------------------ Excel workbook
cols = ["title", "publication_date", "year", "theme",
        "header_label", "header_pos", "header_caut", "header_risk", "header_net",
        "body_label", "body_tokens", "body_pos_per1k", "body_caut_per1k", "body_risk_per1k", "body_net", "consistent", "url"]
with pd.ExcelWriter(f"{OUT}/MIT_AI_framing_results.xlsx", engine="openpyxl") as xw:
    readme = pd.DataFrame({"Item": [
        "Purpose", "Header", "Body", "Categories", "Rule", "Why indexed", "Sensitivity", "Caveat", "Sheets"],
        "Description": [
        "How MIT News language frames AI: positive vs neutral/cautious vs risk-oriented, header vs body.",
        "Title + summary (the line shown under the headline).", "Full article text.",
        "Positive = capability/progress vocabulary; Risk-oriented = challenge/bias/safety/governance vocabulary; Neutral/Cautious = balanced, hedged, or no evaluative wording.",
        "net = (pos_idx - risk_idx)/(pos_idx + risk_idx); >= +0.30 Positive, <= -0.30 Risk-oriented, otherwise Neutral/Cautious.",
        "Positive words are ~4x more frequent than risk words in this corpus, so each rate is divided by its corpus mean (1.0 = typical article) before comparing.",
        "See the Threshold_sensitivity sheet (0.2 to 0.5).",
        "Dictionary method: it counts words, not meaning (negation, irony, topic words such as 'safety' in a safety-engineering story). Validate by reading the Examples sheet.",
        "Article_scores, Category_summary, Header_vs_Body, Vocabulary_rates, Distinctive_body, Distinctive_header, Themes, Examples, Threshold_sensitivity, Lexicon"]})
    readme.to_excel(xw, sheet_name="README", index=False)
    df[cols].to_excel(xw, sheet_name="Article_scores", index=False)
    summ.to_excel(xw, sheet_name="Category_summary", index=False)
    ct.to_excel(xw, sheet_name="Header_vs_Body")
    rates.to_excel(xw, sheet_name="Vocabulary_rates")
    pd.concat({c: dist_body[c] for c in CATS}, names=["category", "rank"]).reset_index(level=0).to_excel(xw, sheet_name="Distinctive_body", index=False)
    pd.concat({c: dist_head[c] for c in CATS}, names=["category", "rank"]).reset_index(level=0).to_excel(xw, sheet_name="Distinctive_header", index=False)
    themes.to_excel(xw, sheet_name="Themes", index=False)
    ex.to_excel(xw, sheet_name="Examples", index=False)
    sens.to_excel(xw, sheet_name="Threshold_sensitivity", index=False)
    lex.to_excel(xw, sheet_name="Lexicon", index=False)
    from openpyxl.styles import Font, PatternFill, Alignment
    for ws in xw.book.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F3B5C")
        for col in ws.columns:
            w = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:60])
            ws.column_dimensions[col[0].column_letter].width = min(max(10, w + 2), 70)
        ws.freeze_panes = "A2"
    xw.book["README"].column_dimensions["B"].width = 120
    for r in xw.book["README"].iter_rows(min_row=2):
        r[1].alignment = Alignment(wrap_text=True, vertical="top")

print("N articles:", len(df))
print(summ.to_string(index=False)); print("\nHeader x Body:\n", ct)
print("\nConsistency (same label):", round(df.consistent.mean(), 3))
print("Header risk -> body positive etc. done"); print(rates)
print("\nSensitivity:\n", sens.to_string(index=False))
print("\nTHEMES\n", themes.drop(columns="top_terms").to_string(index=False))
for c in CATS:
    print(f"\nBODY {c}:", ", ".join(dist_body[c].term[:22]))
for c in CATS:
    print(f"HEADER {c}:", ", ".join(dist_head[c].term[:15]))
